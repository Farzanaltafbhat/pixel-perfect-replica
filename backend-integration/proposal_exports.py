"""Add to the user's existing FastAPI backend; never import into the web app."""
from io import BytesIO
import re
from xml.sax.saxutils import escape
from typing import Literal
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from docx import Document
from docx.shared import Pt
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT

router = APIRouter(prefix="/api/export", tags=["Proposal exports"])

class Section(BaseModel):
    id: str = Field(min_length=1, max_length=100)
    title: str = Field(min_length=1, max_length=300)
    content: str = Field(min_length=1, max_length=100000)

class ProposalExport(BaseModel):
    id: str = Field(min_length=1, max_length=200)
    client_name: str = Field(min_length=1, max_length=200)
    project_title: str = Field(min_length=1, max_length=300)
    document_type: str = Field(min_length=1, max_length=100)
    status: str = Field(max_length=100)
    sections: list[Section] = Field(min_length=1, max_length=30)

def validate_content(proposal: ProposalExport):
    if any(not s.content.strip() for s in proposal.sections):
        raise HTTPException(422, "Complete all empty proposal sections before export.")
    if sum(len(s.content) for s in proposal.sections) > 500000:
        raise HTTPException(413, "Proposal is too large to export.")

def generate_docx(proposal: ProposalExport) -> bytes:
    validate_content(proposal)
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    doc.core_properties.title = proposal.project_title
    doc.core_properties.author = 'BriskBold IT'
    doc.sections[0].header.paragraphs[0].text = 'BriskBold IT · Proposal'
    doc.sections[0].footer.paragraphs[0].text = f'{proposal.id} · {proposal.status} — not an approval record'
    for index, section in enumerate(proposal.sections):
        if index and proposal.sections[index - 1].id == 'cover':
            doc.add_page_break()
        doc.add_heading(section.title, level=1)
        for line in section.content.splitlines():
            doc.add_paragraph(line)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()

def generate_pdf(proposal: ProposalExport) -> bytes:
    validate_content(proposal)
    output = BytesIO()
    styles = getSampleStyleSheet()
    styles['Normal'].alignment = TA_LEFT
    styles['Normal'].wordWrap = 'CJK'
    story = []
    for index, section in enumerate(proposal.sections):
        if index and proposal.sections[index - 1].id == 'cover':
            story.append(PageBreak())
        story.append(Paragraph(escape(section.title), styles['Heading1']))
        for line in section.content.splitlines():
            story.append(Paragraph(escape(line) if line else ' ', styles['Normal']))
            story.append(Spacer(1, 5))
        story.append(Spacer(1, 16))
    def page_brand(canvas, document):
        canvas.saveState()
        canvas.setFont('Helvetica', 9)
        canvas.drawString(40, 810, 'BriskBold IT · Proposal')
        canvas.drawString(40, 25, f'{proposal.status} · Page {document.page}')
        canvas.restoreState()
    SimpleDocTemplate(output, pagesize=(595, 842), leftMargin=40, rightMargin=40, topMargin=55, bottomMargin=45, title=proposal.project_title, author='BriskBold IT').build(story, onFirstPage=page_brand, onLaterPages=page_brand)
    return output.getvalue()

def response(proposal: ProposalExport, format: Literal['docx', 'pdf']):
    content = generate_docx(proposal) if format == 'docx' else generate_pdf(proposal)
    name = re.sub(r'[^a-zA-Z0-9_-]+', '-', f'{proposal.client_name}-{proposal.project_title}')[:100] or 'BriskBold-proposal'
    mime = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document' if format == 'docx' else 'application/pdf'
    return StreamingResponse(BytesIO(content), media_type=mime, headers={'Content-Disposition': f'attachment; filename="{name}.{format}"', 'Cache-Control': 'no-store'})

@router.post('/docx')
def export_docx(proposal: ProposalExport):
    return response(proposal, 'docx')

@router.post('/pdf')
def export_pdf(proposal: ProposalExport):
    return response(proposal, 'pdf')
